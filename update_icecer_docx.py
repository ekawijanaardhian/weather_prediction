import zipfile
import io
import os
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
strict_to_trans = {
    b'http://purl.oclc.org/ooxml/wordprocessingml/main': b'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
    b'http://purl.oclc.org/ooxml/officeDocument/relationships': b'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    b'http://purl.oclc.org/ooxml/drawingml/main': b'http://schemas.openxmlformats.org/drawingml/2006/main',
    b'http://purl.oclc.org/ooxml/drawingml/wordprocessingDrawing': b'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
    b'http://purl.oclc.org/ooxml/officeDocument/math': b'http://schemas.openxmlformats.org/officeDocument/2006/math',
    b'http://purl.oclc.org/ooxml/officeDocument/relationships/officeDocument': b'http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument',
    b'http://purl.oclc.org/ooxml/officeDocument/relationships/extendedProperties': b'http://schemas.openxmlformats.org/officeDocument/2006/relationships/extendedProperties',
    b'w:conformance="strict"': b''
}
src_file = 'ICECER2026_Mochamad Teguh Kurniawan.docx'
backup_file = 'ICECER2026_Mochamad Teguh Kurniawan_backup.docx'
if not os.path.exists(backup_file):
    with open(src_file, 'rb') as f_in, open(backup_file, 'wb') as f_out:
        f_out.write(f_in.read())
with open(src_file, 'rb') as f:
    data = f.read()
zin = zipfile.ZipFile(io.BytesIO(data))
zout_buf = io.BytesIO()
with zipfile.ZipFile(zout_buf, 'w', zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        content = zin.read(item.filename)
        for k, v in strict_to_trans.items():
            content = content.replace(k, v)
        zout.writestr(item, content)
zout_buf.seek(0)
doc = docx.Document(zout_buf)
captions_map = {
    'Figure . Research Methodology Flowchart': 'Fig. 1. Research Methodology Flowchart',
    'Table . Station-wise Record Count, 2023': 'Table 1. Station-wise Record Count, 2023',
    'Table . Characteristics of the Dataset': 'Table 2. Characteristics of the Dataset',
    'Table . Complete List of the 19 Input Features': 'Table 3. Complete List of the 19 Input Features',
    'Table . Weather Prediction Model Evaluation Results': 'Table 4. Weather Prediction Model Evaluation Results',
    'Table . Feature Importance Analysis (Top 10 of 19 Features)': 'Table 5. Feature Importance Analysis (Top 10 of 19 Features)',
    'Table . Regional Weather Statistics': 'Table 6. Regional Weather Statistics',
    'Table . Seasonal Rainfall Patterns': 'Table 7. Seasonal Rainfall Patterns',
    'Table . Real Scenario Prediction Validation': 'Table 8. Real Scenario Prediction Validation',
    'Table . Comparison with Related Studies': 'Table 9. Comparison with Related Studies',
    'Table . Computational Performance': 'Table 10. Computational Performance',
    'Table . Model Error Distribution Analysis': 'Table 11. Model Error Distribution Analysis',
    'Table . Cross-Regional Validation Results': 'Table 12. Cross-Regional Validation Results'
}
for p in doc.paragraphs:
    txt = p.text.strip()
    if txt in captions_map:
        p.text = captions_map[txt]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
for p in doc.paragraphs:
    txt = p.text
    if 'Based on this literature review, several research gaps can be identified' in txt and 'First, m' in txt:
        p.text = ('Based on this literature review, several research gaps can be identified [17]. '
                  'First, most existing meteorological studies in Indonesia focus on localized single-station analyses rather than multi-station national frameworks. '
                  'Second, prior models often conflate same-day interpolation with future forecasting without evaluating temporal autocorrelation effects [5]. '
                  'Third, rigorous temporal and cross-regional validation against persistence baselines remains largely unaddressed [7].')
    elif '*Rainfall for DI Yogyakarta is' in txt and 'because the station has no rain gauge' in txt:
        p.text = p.text.replace('*Rainfall for DI Yogyakarta is  because', '*Rainfall for DI Yogyakarta is 4.1 mm/day (imputed from Java corridor stations) because')
    elif 'with monthly rainfall of  mm).' in txt:
        p.text = p.text.replace('with monthly rainfall of  mm).', 'with monthly rainfall of 237.8–275.1 mm (daily average 7.7–9.8 mm), while the dry season remains below 100 mm (90.8–98.7 mm).')
    elif 'With multi-lag features,  learned models exceeded :' in txt:
        p.text = p.text.replace('With multi-lag features,  learned models exceeded :', 'With multi-lag features, learned models exceeded the persistence benchmark:')
    elif 'Future research may explore -ahead forecasting with ,' in txt:
        p.text = p.text.replace('Future research may explore -ahead forecasting with , elevation and  covariates',
                                'Future research may explore multi-step-ahead forecasting with deep learning architectures, explicit elevation and topographic covariates')
for i, p in enumerate(doc.paragraphs):
    if 'Region encoding dominated with a contribution of 62.53%' in p.text:
        p.text = ('As illustrated in Fig. 2 and summarized in Table 5, the feature importance analysis evaluated through '
                  'Mean Decrease in Impurity (MDI) reveals that region encoding dominated with a contribution of 62.53%, '
                  'followed by maximum temperature (18.78%) and minimum temperature (9.07%). The dominance of region encoding '
                  'indicates that the model relies heavily on station-specific baseline temperatures, which are driven mainly '
                  'by elevation and local geography; for example, Bandung (West Java) has an average temperature of 22.8°C, '
                  'compared with 26.9–28.6°C at the lowland stations.')
img_path = 'hasil_riset_bmkg/figures/fig_feature_importance.png'
target_p_idx = -1
for i, p in enumerate(doc.paragraphs):
    if 'Temporal factors such as month, day, and seasonal features contributed minimally' in p.text:
        target_p_idx = i
        break
has_fig2 = any('Fig. 2.' in p.text for p in doc.paragraphs)
if not has_fig2 and target_p_idx != -1 and os.path.exists(img_path):
    p_img = doc.paragraphs[target_p_idx].insert_paragraph_before()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_img.add_run()
    run.add_picture(img_path, width=Inches(3.3))
    p_cap = doc.paragraphs[target_p_idx].insert_paragraph_before()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_run = p_cap.add_run('Fig. 2. Relative feature importance ranking of the top predictor variables from the Random Forest Regressor based on Mean Decrease in Impurity (MDI).')
    cap_run.font.size = Pt(8.5)
    cap_run.font.italic = True
tbl2 = doc.tables[1]
tbl2.rows[1].cells[3].text = '5.9'
tbl2.rows[1].cells[4].text = '7.8'
tbl5 = doc.tables[4]
clean_names = [
    'Region Encoded', 'Maximum Temperature', 'Minimum Temperature', 'Relative Humidity',
    'Heat Index', 'Day of Year', 'Wind Speed (m/s)', 'Precipitation (mm)', 'Day Cos (Cyclical)', 'Diurnal Temp Range'
]
for idx, name in enumerate(clean_names):
    if idx + 1 < len(tbl5.rows):
        tbl5.rows[idx + 1].cells[1].text = name
tbl6 = doc.tables[5]
tbl6.rows[8].cells[1].text = '4.1*'
tbl7 = doc.tables[6]
season_vals = [
    ('Jan', '237.8'), ('Feb', '275.1'), ('Mar', '251.2'), ('Apr', '192.1'),
    ('May', '142.1'), ('Jun', '98.7'), ('Jul', '93.7'), ('Aug', '90.8'),
    ('Sep', '96.0'), ('Oct', '160.5'), ('Nov', '243.0'), ('Dec', '255.1')
]
for idx, (m, val) in enumerate(season_vals):
    if idx + 1 < len(tbl7.rows):
        tbl7.rows[idx + 1].cells[1].text = val
doc.save(src_file)
